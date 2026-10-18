from io import BytesIO
from zipfile import ZipFile

import fitz
from docx import Document

from app.core.config import get_settings


class DocumentValidationError(ValueError):
    pass


def extract_job_text(filename: str, content_type: str | None, data: bytes) -> str:
    settings = get_settings()
    if not data or len(data) > settings.max_job_upload_bytes:
        raise DocumentValidationError("The uploaded file is empty or exceeds the upload size limit.")
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in {"pdf", "docx"}:
        raise DocumentValidationError("Upload a PDF or DOCX file.")
    expected_type = "application/pdf" if suffix == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if content_type and content_type not in {"application/octet-stream", expected_type}:
        raise DocumentValidationError("The uploaded file type does not match PDF or DOCX.")

    try:
        if suffix == "pdf":
            if not data.startswith(b"%PDF-"):
                raise DocumentValidationError("The file is not a valid PDF.")
            with fitz.open(stream=data, filetype="pdf") as pdf:
                if pdf.page_count > settings.max_job_pdf_pages:
                    raise DocumentValidationError("The PDF has too many pages to analyze.")
                page_text = []
                total_chars = 0
                for page in pdf:
                    page_content = page.get_text()
                    page_text.append(page_content)
                    total_chars += len(page_content)
                    if total_chars > settings.max_job_text_chars:
                        raise DocumentValidationError("The extracted job description is too long to analyze.")
                text = "\n".join(page_text)
        else:
            if not data.startswith(b"PK"):
                raise DocumentValidationError("The file is not a valid DOCX document.")
            with ZipFile(BytesIO(data)) as archive:
                if "word/document.xml" not in archive.namelist() or len(archive.namelist()) > 2000:
                    raise DocumentValidationError("The DOCX document structure is invalid.")
                if sum(item.file_size for item in archive.infolist()) > settings.max_job_upload_bytes * 20:
                    raise DocumentValidationError("The DOCX document expands beyond the allowed size.")
            document = Document(BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            text = "\n".join(parts)
    except DocumentValidationError:
        raise
    except Exception as error:
        raise DocumentValidationError("The document could not be read. Please upload a valid PDF or DOCX.") from error

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip()).strip()
    if len(text) < 40:
        raise DocumentValidationError("No readable job description text was found in the document.")
    if len(text) > settings.max_job_text_chars:
        raise DocumentValidationError("The extracted job description is too long to analyze.")
    return text
